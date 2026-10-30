"""PowerPoint status deck built with python-pptx.

Charts are native PowerPoint charts (with embedded data), so they stay
editable in PowerPoint: right-click -> Edit Data shows the numbers. The Gantt
and heat map are drawn from shapes and tables. Every slide carries speaker
notes that say where its numbers come from.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .analysis import Analysis, rag
from .model import risk_rating

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT_GREY = RGBColor(0xD9, 0xD9, 0xD9)
BLUE = RGBColor(0x44, 0x72, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RAG_COLOR = {"Green": RGBColor(0x70, 0xAD, 0x47), "Amber": RGBColor(0xFF, 0xC0, 0x00),
             "Red": RGBColor(0xC0, 0x00, 0x00), "N/A": RGBColor(0xA6, 0xA6, 0xA6)}
LEVEL_COLOR = {"High": RGBColor(0xF8, 0xCB, 0xAD), "Medium": RGBColor(0xFF, 0xE6, 0x99),
               "Low": RGBColor(0xC6, 0xE0, 0xB4)}
GANTT_ROWS = 28
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def forecast_rag(baseline: float, p80: float) -> str:
    """Slip of the P80 forecast relative to baseline duration."""
    if baseline <= 0:
        return "N/A"
    slip = (p80 - baseline) / baseline
    if slip <= 0.05:
        return "Green"
    if slip <= 0.15:
        return "Amber"
    return "Red"


def risk_rag(high_risks: int) -> str:
    return "Green" if high_risks == 0 else "Amber" if high_risks <= 2 else "Red"


def _money(v: float | None) -> str:
    if v is None:
        return "n/a"
    sign = "-" if v < 0 else ""
    v = abs(v)
    if v >= 1e6:
        return f"{sign}${v / 1e6:.2f}M"
    if v >= 1e3:
        return f"{sign}${v / 1e3:.1f}k"
    return f"{sign}${v:.0f}"


def _idx(v: float | None) -> str:
    return "n/a" if v is None else f"{v:.2f}"


def _d(d: date) -> str:
    return d.strftime("%d %b %Y")


class _Deck:
    def __init__(self, a: Analysis):
        self.a = a
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.blank = self.prs.slide_layouts[6]

    # ----------------------------------------------------------- primitives
    def slide(self, title: str, notes: str):
        s = self.prs.slides.add_slide(self.blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        self.text(s, title, Inches(0.4), Inches(0.12), Inches(11), Inches(0.7), size=26,
                  bold=True, color=WHITE)
        self.text(s, f"{self.a.program.name} | status {_d(self.a.program.status_date)}",
                  Inches(0.4), Inches(7.05), Inches(9), Inches(0.35), size=10, color=GREY)
        s.notes_slide.notes_text_frame.text = notes
        return s

    @staticmethod
    def text(s, text, x, y, w, h, size=14, bold=False, color=GREY, align=None):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        lines = text if isinstance(text, list) else [text]
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
            if align:
                p.alignment = align
        return tb

    @staticmethod
    def table(s, rows, x, y, w, col_widths, size=11):
        t = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.3 * len(rows))).table
        for j, cw in enumerate(col_widths):
            t.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = t.cell(i, j)
                cell.text = str(v)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(size)
                para.font.bold = i == 0
        return t

    # ---------------------------------------------------------------- slides
    def title_slide(self):
        a = self.a
        s = self.prs.slides.add_slide(self.blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY
        bg.line.fill.background()
        self.text(s, a.program.name, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.2),
                  size=44, bold=True, color=WHITE)
        self.text(s, f"Program status report - data as of {_d(a.program.status_date)}",
                  Inches(0.8), Inches(3.6), Inches(11.5), Inches(0.6), size=22, color=WHITE)
        self.text(s, f"{len(a.program.tasks)} tasks | {len(a.program.risks)} risks | "
                     f"{a.mc.iterations:,} Monte Carlo iterations",
                  Inches(0.8), Inches(4.3), Inches(11.5), Inches(0.5), size=16, color=LIGHT_GREY)
        s.notes_slide.notes_text_frame.text = "Generated by statusdeck from the program tracker."

    def summary_slide(self):
        a, e, mc = self.a, self.a.evm, self.a.mc
        p80 = mc.percentile(80)
        high = sum(1 for r in a.program.risks if r.rating == "High")
        tiles = [
            ("Schedule (SPI)", _idx(e.spi), rag(e.spi)),
            ("Cost (CPI)", _idx(e.cpi), rag(e.cpi)),
            ("P80 finish", _d(a.date_at(p80)), forecast_rag(a.cpm.finish, p80)),
            ("High risks", str(high), risk_rag(high)),
        ]
        s = self.slide("Executive summary",
                       "SPI/CPI from the EVM sheet; P80 from the MonteCarlo sheet; risk counts "
                       "from the Risks sheet. RAG: index >= 0.95 green, >= 0.85 amber; P80 slip "
                       "<= 5% green, <= 15% amber; 0 high risks green, <= 2 amber.")
        for k, (label, value, status) in enumerate(tiles):
            x = Inches(0.5 + k * 3.15)
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.2), Inches(2.95),
                                     Inches(1.6))
            box.fill.solid()
            box.fill.fore_color.rgb = RAG_COLOR[status]
            box.line.fill.background()
            box.name = f"KPI {label}"
            tf = box.text_frame
            tf.text = label
            tf.paragraphs[0].font.size = Pt(14)
            p = tf.add_paragraph()
            p.text = value
            p.font.size, p.font.bold = Pt(28), True
            p = tf.add_paragraph()
            p.text = status
            p.font.size = Pt(12)
            for para in tf.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                para.font.color.rgb = WHITE if status in ("Red", "Green") else NAVY
        crit = a.cpm.critical_ids
        top = int(mc.criticality.argmax())
        bullets = [
            f"Schedule: EV {_money(e.ev)} against PV {_money(e.pv)} "
            f"(SV {_money(e.sv)}, SPI {_idx(e.spi)}).",
            f"Cost: AC {_money(e.ac)}, CPI {_idx(e.cpi)}; estimate at completion "
            f"{_money(e.eac)} vs budget {_money(e.bac)} (VAC {_money(e.vac)}).",
            f"Baseline finish {_d(a.baseline_finish_date)}; Monte Carlo P50 "
            f"{_d(a.date_at(mc.percentile(50)))}, P80 {_d(a.date_at(p80))}; "
            f"chance of meeting baseline {mc.p_on_time:.0%}.",
            f"{len(crit)} of {len(a.program.tasks)} tasks have zero float; most frequently "
            f"critical in simulation: {a.program.tasks[top].id} {a.program.tasks[top].name} "
            f"({mc.criticality[top]:.0%} of iterations).",
            f"Risks: {high} high, "
            f"{sum(1 for r in a.program.risks if r.rating == 'Medium')} medium, "
            f"{sum(1 for r in a.program.risks if r.rating == 'Low')} low.",
        ]
        self.text(s, ["- " + b for b in bullets], Inches(0.5), Inches(3.1), Inches(12.3),
                  Inches(3.8), size=17, color=NAVY)

    def gantt_slide(self):
        a, cpm = self.a, self.a.cpm
        n = len(a.program.tasks)
        chosen = sorted(range(n), key=lambda i: (cpm.total_float[i], cpm.es[i]))[:GANTT_ROWS]
        chosen.sort(key=lambda i: (cpm.es[i], cpm.ef[i]))
        shown = "all tasks" if n <= GANTT_ROWS else f"the {GANTT_ROWS} lowest-float tasks"
        s = self.slide("Schedule - critical path",
                       f"Shows {shown} from the Schedule sheet, ordered by early start. Red = "
                       "zero total float (critical). Dashed line = status date.")
        left, label_w, top = Inches(0.4), Inches(3.4), Inches(1.3)
        x0 = left + label_w
        width = SLIDE_W - x0 - Inches(0.4)
        horizon = max(cpm.finish, a.evm.elapsed, 1)
        row_h = Emu(int((Inches(5.55)) / max(len(chosen), 1)))
        scale = width / horizon

        def x(t):
            return x0 + Emu(int(t * scale))

        # month gridlines
        last_month = None
        for k in range(int(horizon) + 1):
            d = a.calendar.date_of(k)
            if (d.year, d.month) != last_month:
                last_month = (d.year, d.month)
                ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x(k), top - Inches(0.1), x(k),
                                            top + row_h * len(chosen))
                ln.line.color.rgb = LIGHT_GREY
                self.text(s, d.strftime("%b %y"), min(x(k), SLIDE_W - Inches(0.9)), Inches(0.95),
                          Inches(0.9), Inches(0.3), size=9)
        font = max(7, min(11, int(row_h / 12700 * 0.55)))
        for row, i in enumerate(chosen):
            t = a.program.tasks[i]
            y = top + row_h * row
            self.text(s, f"{t.id} {t.name}"[:44], left, y - Inches(0.06), label_w, row_h,
                      size=font, color=NAVY)
            crit = cpm.is_critical(i)
            color = RAG_COLOR["Red"] if crit else BLUE
            if cpm.duration[i] == 0:
                d = Emu(int(row_h * 0.7))
                shp = s.shapes.add_shape(MSO_SHAPE.DIAMOND, x(cpm.es[i]) - d // 2,
                                         y + Emu(int(row_h * 0.15)), d, d)
            else:
                shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x(cpm.es[i]),
                                         y + Emu(int(row_h * 0.15)),
                                         max(Emu(int(cpm.duration[i] * scale)), Emu(12700)),
                                         Emu(int(row_h * 0.7)))
            shp.fill.solid()
            shp.fill.fore_color.rgb = color
            shp.line.fill.background()
            shp.name = f"bar {t.id}"
        sx = x(a.evm.elapsed)
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sx, top - Inches(0.2), sx,
                                    top + row_h * len(chosen))
        ln.line.color.rgb = GREY
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln.name = "status date"

    def critical_path_slide(self):
        a, cpm, mc = self.a, self.a.cpm, self.a.mc
        path = cpm.driving_path()
        s = self.slide("Driving path",
                       "The zero-float chain ending at the latest-finishing task, from the "
                       "Schedule sheet. MC criticality = share of simulated iterations in which "
                       "the task had zero float.")
        rows = [["ID", "Task", "Owner", "Start", "Finish", "Days", "MC criticality"]]
        for tid in path[:16]:
            i = cpm.network.index[tid]
            t = a.program.tasks[i]
            rows.append([t.id, t.name, t.owner, _d(a.task_start(i)), _d(a.task_finish(i)),
                         f"{cpm.duration[i]:g}", f"{mc.criticality[i]:.0%}"])
        self.table(s, rows, Inches(0.5), Inches(1.2), Inches(12.3),
                   [0.8, 4.2, 1.6, 1.6, 1.6, 0.9, 1.6], size=11)
        if len(path) > 16:
            self.text(s, f"(+{len(path) - 16} more tasks on the path)", Inches(0.5), Inches(6.6),
                      Inches(6), Inches(0.4), size=11)

    def evm_slide(self):
        a, e = self.a, self.a.evm
        s = self.slide("Earned value",
                       "Planned value curve = sum of budget x straight-line planned % at each "
                       "week (EVM sheet). EV and AC are plotted at the status date.")
        status_k = min(range(len(e.curve)), key=lambda k: abs(e.curve[k][0] - e.elapsed))
        cd = CategoryChartData()
        cd.categories = [d.strftime("%d %b") for _, d, _ in e.curve]
        cd.add_series("Planned value", [round(v, 2) for _, _, v in e.curve])
        cd.add_series("Earned value", [round(e.ev, 2) if k == status_k else None
                                       for k in range(len(e.curve))])
        cd.add_series("Actual cost", [round(e.ac, 2) if k == status_k else None
                                      for k in range(len(e.curve))])
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.4), Inches(1.1), Inches(8.6),
                                Inches(5.8), cd)
        ch = gf.chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.value_axis.tick_labels.number_format = '"$"#,##0'
        ch.value_axis.tick_labels.number_format_is_linked = False
        ch.category_axis.tick_labels.font.size = Pt(9)
        pv = ch.plots[0].series[0]
        pv.marker.style = XL_MARKER_STYLE.NONE
        pv.smooth = False
        for sr in list(ch.plots[0].series)[1:]:
            sr.marker.style = XL_MARKER_STYLE.CIRCLE
            sr.marker.size = 11
        rows = [["Metric", "Value"], ["BAC", _money(e.bac)], ["PV", _money(e.pv)],
                ["EV", _money(e.ev)], ["AC", _money(e.ac)], ["SV", _money(e.sv)],
                ["CV", _money(e.cv)], ["SPI", _idx(e.spi)], ["CPI", _idx(e.cpi)],
                ["EAC", _money(e.eac)], ["TCPI", _idx(e.tcpi)]]
        self.table(s, rows, Inches(9.3), Inches(1.2), Inches(3.6), [1.6, 2.0], size=12)

    def risk_sim_slide(self):
        a, mc = self.a, self.a.mc
        counts, edges = mc.histogram(20)
        s = self.slide("Schedule risk (Monte Carlo)",
                       f"{mc.iterations:,} iterations; each task duration drawn from "
                       "triangular(optimistic, most likely, pessimistic). Bars = iterations "
                       "finishing in each working-day bin (MonteCarlo sheet).")
        cd = CategoryChartData()
        cd.categories = [f"{edges[k]:.0f}-{edges[k + 1]:.0f}" for k in range(len(counts))]
        cd.add_series("Iterations", [int(c) for c in counts])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.4), Inches(1.1),
                                Inches(8.6), Inches(5.8), cd)
        ch = gf.chart
        ch.has_legend = False
        ch.plots[0].gap_width = 15
        ch.category_axis.tick_labels.font.size = Pt(9)
        ch.category_axis.has_title = True
        ch.category_axis.axis_title.text_frame.text = "Project finish (working days)"
        rows = [["Measure", "Working days", "Date"],
                ["Plan", f"{a.cpm.finish:.0f}", _d(a.baseline_finish_date)]]
        for q, v in mc.percentiles.items():
            if q in (10, 50, 80, 90):
                rows.append([f"P{q}", f"{v:.1f}", _d(a.date_at(v))])
        rows.append(["P(<= plan)", f"{mc.p_on_time:.1%}", ""])
        self.table(s, rows, Inches(9.3), Inches(1.2), Inches(3.6), [1.1, 1.1, 1.4], size=12)
        self.text(s, "Merge bias: parallel paths that each look on time push the joint "
                     "finish later than the single most-likely plan.",
                  Inches(9.3), Inches(4.4), Inches(3.6), Inches(1.5), size=11)

    def criticality_slide(self):
        a, mc = self.a, self.a.mc
        top = sorted(range(len(a.program.tasks)), key=lambda i: -mc.criticality[i])[:10]
        top.reverse()  # bar charts plot bottom-up
        s = self.slide("Schedule drivers",
                       "Criticality index: fraction of Monte Carlo iterations in which the task "
                       "had zero total float. Focus mitigation on the top of this list.")
        cd = CategoryChartData()
        cd.categories = [f"{a.program.tasks[i].id} {a.program.tasks[i].name}"[:40] for i in top]
        cd.add_series("Criticality index", [round(float(mc.criticality[i]), 4) for i in top])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.4), Inches(1.1),
                                Inches(12.4), Inches(5.8), cd)
        ch = gf.chart
        ch.has_legend = False
        ch.value_axis.maximum_scale = 1.0
        ch.value_axis.tick_labels.number_format = "0%"
        ch.value_axis.tick_labels.number_format_is_linked = False
        plot = ch.plots[0]
        plot.has_data_labels = True
        plot.data_labels.number_format = "0%"
        plot.data_labels.number_format_is_linked = False

    def risk_slide(self):
        a = self.a
        risks = a.program.risks
        s = self.slide("Risk register",
                       "Heat map counts risks by probability (rows) and impact (columns), "
                       "from the Risks sheet. Score = probability x impact; >= 15 high, "
                       ">= 8 medium.")
        t = s.shapes.add_table(6, 6, Inches(0.5), Inches(1.4), Inches(4.6), Inches(3.6)).table
        t.cell(0, 0).text = "P \\ I"
        for i in range(1, 6):
            t.cell(0, i).text = str(i)
        for row, p in enumerate(range(5, 0, -1), start=1):
            t.cell(row, 0).text = str(p)
            for i in range(1, 6):
                c = t.cell(row, i)
                n = sum(1 for r in risks if r.probability == p and r.impact == i)
                c.text = str(n) if n else ""
                c.fill.solid()
                c.fill.fore_color.rgb = LEVEL_COLOR[risk_rating(p * i)]
                c.text_frame.paragraphs[0].font.color.rgb = NAVY
                c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        self.text(s, "Probability (rows) x Impact (columns)", Inches(0.5), Inches(5.1),
                  Inches(4.6), Inches(0.4), size=11)
        top = sorted(risks, key=lambda r: (-r.score, r.id))[:8]
        rows = [["ID", "Risk", "Score", "Owner", "Mitigation"]]
        rows += [[r.id, r.title, f"{r.score} {r.rating}", r.owner, r.mitigation] for r in top]
        self.table(s, rows, Inches(5.4), Inches(1.4), Inches(7.5), [0.5, 2.6, 1.0, 1.1, 2.3],
                   size=10)

    def build(self):
        self.title_slide()
        self.summary_slide()
        self.gantt_slide()
        self.critical_path_slide()
        self.evm_slide()
        self.risk_sim_slide()
        self.criticality_slide()
        self.risk_slide()
        return self.prs


def build_deck(analysis: Analysis, path: str | Path | None = None):
    prs = _Deck(analysis).build()
    if path is not None:
        prs.save(path)
    return prs
