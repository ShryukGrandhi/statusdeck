"""Read/write the input tracker workbook (sheets: Program, Tasks, Risks, Holidays).

Columns are located by header text (case-insensitive), so users may reorder
or add columns in their tracker without breaking the import.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .model import Program, Risk, Task, parse_preds

TASK_COLS = {
    "id": "ID", "name": "Task", "workstream": "Workstream", "owner": "Owner",
    "preds": "Predecessors", "optimistic": "Optimistic (d)", "likely": "Most Likely (d)",
    "pessimistic": "Pessimistic (d)", "budget": "Budget", "pct_complete": "% Complete",
    "actual_cost": "Actual Cost",
}
RISK_COLS = {
    "id": "ID", "title": "Risk", "owner": "Owner", "probability": "Probability (1-5)",
    "impact": "Impact (1-5)", "mitigation": "Mitigation",
}
_ALIASES = {
    "name": {"task", "name", "task name"},
    "preds": {"predecessors", "depends on", "preds"},
    "optimistic": {"optimistic (d)", "optimistic"},
    "likely": {"most likely (d)", "most likely", "likely", "duration"},
    "pessimistic": {"pessimistic (d)", "pessimistic"},
    "pct_complete": {"% complete", "percent complete", "pct complete"},
    "actual_cost": {"actual cost", "ac"},
    "title": {"risk", "title", "description"},
    "probability": {"probability (1-5)", "probability"},
    "impact": {"impact (1-5)", "impact"},
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _as_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v).strip())


def _num(v, default=0.0) -> float:
    if v is None or (isinstance(v, str) and not v.strip()):
        return default
    if isinstance(v, str):
        v = v.strip()
        if v.endswith("%"):
            return float(v[:-1]) / 100
        return float(v.replace(",", "").replace("$", ""))
    return float(v)


def _column_map(header_row, spec: dict[str, str]) -> dict[str, int]:
    cols: dict[str, int] = {}
    norm = [str(h).strip().lower() if h is not None else "" for h in header_row]
    for key, label in spec.items():
        names = _ALIASES.get(key, set()) | {label.lower(), key}
        for j, h in enumerate(norm):
            if h in names:
                cols[key] = j
                break
    return cols


def _rows(ws, spec, required):
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return
    cols = _column_map(header, spec)
    missing = [spec[k] for k in required if k not in cols]
    if missing:
        raise ValueError(f"sheet {ws.title!r} missing column(s): {', '.join(missing)}")
    for row in rows:
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        yield {k: (row[j] if j < len(row) else None) for k, j in cols.items()}


def read_tracker(path: str | Path) -> Program:
    wb = load_workbook(path, data_only=True)
    info = {}
    for row in wb["Program"].iter_rows(values_only=True):
        if row and row[0] is not None:
            info[str(row[0]).strip().lower()] = row[1]
    tasks = []
    for r in _rows(wb["Tasks"], TASK_COLS, ["id", "name", "likely"]):
        likely = _num(r.get("likely"))
        pct = _num(r.get("pct_complete"))
        if pct > 1.0:  # tolerate 0-100 entry
            pct /= 100.0
        tasks.append(Task(
            id=str(r["id"]).strip(),
            name=str(r["name"] or "").strip(),
            workstream=str(r.get("workstream") or "").strip(),
            owner=str(r.get("owner") or "").strip(),
            preds=parse_preds(r.get("preds")),
            optimistic=_num(r.get("optimistic"), likely),
            likely=likely,
            pessimistic=_num(r.get("pessimistic"), likely),
            budget=_num(r.get("budget")),
            pct_complete=pct,
            actual_cost=_num(r.get("actual_cost")),
        ))
    risks = []
    if "Risks" in wb.sheetnames:
        for r in _rows(wb["Risks"], RISK_COLS, ["id", "title", "probability", "impact"]):
            risks.append(Risk(
                id=str(r["id"]).strip(), title=str(r["title"]).strip(),
                owner=str(r.get("owner") or "").strip(),
                probability=int(_num(r["probability"])), impact=int(_num(r["impact"])),
                mitigation=str(r.get("mitigation") or "").strip(),
            ))
    holidays = []
    if "Holidays" in wb.sheetnames:
        for row in wb["Holidays"].iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                holidays.append(_as_date(row[0]))
    return Program(
        name=str(info.get("program", "Program")),
        start=_as_date(info["start date"]),
        status_date=_as_date(info["status date"]),
        tasks=tasks, risks=risks, holidays=sorted(holidays),
    )


def _header(ws, labels):
    ws.append(labels)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    ws.freeze_panes = "A2"


def write_tracker(program: Program, path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Program"
    ws.append(["Program", program.name])
    ws.append(["Start Date", program.start])
    ws.append(["Status Date", program.status_date])
    for r in (2, 3):
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40

    ws = wb.create_sheet("Tasks")
    _header(ws, list(TASK_COLS.values()))
    for t in program.tasks:
        ws.append([t.id, t.name, t.workstream, t.owner, t.preds_text(), t.optimistic,
                   t.likely, t.pessimistic, t.budget, t.pct_complete, t.actual_cost])
    n = len(program.tasks) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        row[8].number_format = row[10].number_format = '"$"#,##0'
        row[9].number_format = "0%"
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                        showErrorMessage=True, error="Enter a percentage between 0% and 100%")
    dv.add(f"J2:J{max(n, 2)}")
    ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGHIJK", (7, 40, 14, 14, 18, 13, 15, 14, 12, 11, 12)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Risks")
    _header(ws, list(RISK_COLS.values()))
    for r in program.risks:
        ws.append([r.id, r.title, r.owner, r.probability, r.impact, r.mitigation])
    for col, w in zip("ABCDEF", (7, 48, 14, 16, 13, 48)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Holidays")
    _header(ws, ["Holiday"])
    for h in program.holidays:
        ws.append([h])
        ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 14
    wb.save(path)
