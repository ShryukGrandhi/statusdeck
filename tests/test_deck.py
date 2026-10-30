import pytest
from pptx import Presentation
from pptx.util import Emu

from statusdeck.deck import build_deck, forecast_rag, risk_rag


@pytest.fixture(scope="module")
def deck(tmp_path_factory, sample_analysis):
    path = tmp_path_factory.mktemp("d") / "d.pptx"
    build_deck(sample_analysis, path)
    return Presentation(path)


def _titles(prs):
    out = []
    for s in prs.slides:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text]
        out.append(texts[0] if texts else "")
    return out


def test_slide_order(deck):
    assert _titles(deck)[1:] == ["Executive summary", "Schedule - critical path", "Driving path",
                                 "Earned value", "Schedule risk (Monte Carlo)",
                                 "Schedule drivers", "Risk register"]


def test_every_slide_has_notes_and_stays_on_canvas(deck):
    for s in deck.slides:
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= deck.slide_width + Emu(1)
            assert sh.top + sh.height <= deck.slide_height + Emu(1)


def test_kpi_tiles_match_analysis(deck, sample_analysis):
    e = sample_analysis.evm
    tiles = {sh.name: sh.text_frame.text for sh in deck.slides[1].shapes
             if sh.name.startswith("KPI")}
    assert tiles["KPI Schedule (SPI)"].split("\n")[1] == f"{e.spi:.2f}"
    assert tiles["KPI Cost (CPI)"].split("\n")[1] == f"{e.cpi:.2f}"


def test_charts_carry_engine_numbers(deck, sample_analysis):
    a = sample_analysis
    charts = [sh.chart for s in deck.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 3
    s_curve, hist, crit = charts
    pv = list(s_curve.plots[0].series[0].values)
    assert pv[-1] == pytest.approx(a.evm.bac)
    ev = [v for v in s_curve.plots[0].series[1].values if v is not None]
    assert ev == [pytest.approx(a.evm.ev)]
    assert sum(hist.plots[0].series[0].values) == a.mc.iterations
    vals = list(crit.plots[0].series[0].values)
    assert vals[-1] == pytest.approx(float(a.mc.criticality.max()), abs=1e-4)


def test_gantt_marks_critical_tasks_red(deck, sample_analysis):
    a = sample_analysis
    bars = {sh.name[4:]: sh for sh in deck.slides[2].shapes if sh.name.startswith("bar ")}
    assert 0 < len(bars) <= 28
    for tid, sh in bars.items():
        i = a.cpm.network.index[tid]
        red = str(sh.fill.fore_color.rgb) == "C00000"
        assert red == a.cpm.is_critical(i)


def test_driving_path_table(deck, sample_analysis):
    tables = [sh.table for sh in deck.slides[3].shapes if sh.has_table]
    ids = [tables[0].cell(r, 0).text for r in range(1, len(tables[0].rows))]
    assert ids == sample_analysis.cpm.driving_path()[:16]


def test_rag_helpers():
    assert forecast_rag(100, 104) == "Green"
    assert forecast_rag(100, 114) == "Amber"
    assert forecast_rag(100, 130) == "Red"
    assert forecast_rag(0, 5) == "N/A"
    assert [risk_rag(n) for n in (0, 2, 3)] == ["Green", "Amber", "Red"]
