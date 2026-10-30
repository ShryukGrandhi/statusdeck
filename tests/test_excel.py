"""The workbook's formulas are recalculated by pycel (an independent Excel formula
engine) and must reproduce the Python engine's numbers cell for cell."""

from datetime import date

import pytest
from openpyxl import load_workbook

from statusdeck.analysis import analyze
from statusdeck.excel_report import build_workbook
from statusdeck.sample import generate
from statusdeck.verify import _compat, verify_workbook


@pytest.fixture(scope="module")
def built(tmp_path_factory, sample_analysis):
    path = tmp_path_factory.mktemp("x") / "w.xlsx"
    _, expected = build_workbook(sample_analysis, path)
    return path, expected


def test_every_formula_cell_recalculates_to_engine_value(built):
    path, expected = built
    rep = verify_workbook(path, expected)
    assert rep.checked > 1000
    assert rep.ok, rep.mismatches[:5]


def test_formulas_are_live_not_pasted(built):
    path, _ = built
    wb = load_workbook(path)
    s = wb["Schedule"]
    assert s["J3"].value == "=MAX(0)"
    assert str(s["K3"].value).startswith("=") and str(s["M3"].value).startswith("=MIN(ProjFinish")
    assert wb["EVM"]["F3"].value.startswith("=IF(")
    assert set(wb.defined_names) >= {"CalDates", "StatusDate", "Elapsed", "ProjFinish"}
    assert wb.calculation.fullCalcOnLoad


def test_sheets_charts_and_formatting(built):
    path, _ = built
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Schedule", "EVM", "Risks", "MonteCarlo", "Calendar"]
    assert len(wb["Schedule"].conditional_formatting) >= 3  # Gantt bars + critical flag
    assert wb["Risks"].data_validations.dataValidation
    assert wb["EVM"].data_validations.dataValidation
    # openpyxl does not read charts back, so check the package parts instead
    import zipfile
    names = zipfile.ZipFile(path).namelist()
    assert sum(n.startswith("xl/charts/chart") for n in names) == 2


def _recalc_after_edit(path, edits, outputs):
    """Edit input cells in the saved workbook and recalculate with pycel."""
    _compat()
    from pycel import ExcelCompiler
    xl = ExcelCompiler(filename=str(path))
    for ref in outputs:  # pycel only knows cells already in its dependency graph
        xl.evaluate(ref)
    for ref, v in edits.items():
        xl.set_value(ref, v)
    return xl


def test_editing_a_duration_replans_downstream(small_program, tmp_path):
    a = analyze(small_program, iterations=200)
    path = tmp_path / "s.xlsx"
    build_workbook(a, path)
    # Task C is row 5 (A=3, B=4, C=5). C has 1 day float; +3 days makes it critical.
    xl = _recalc_after_edit(path, {"Schedule!G5": 5}, ["Summary!B6", "Schedule!P4", "Schedule!P5"])
    assert xl.evaluate("Summary!B6") == 16  # project finish 14 -> 16
    assert xl.evaluate("Schedule!P5") == "Yes"
    assert xl.evaluate("Schedule!P4") == "No"  # B now has float


def test_editing_status_date_moves_earned_value(small_program, tmp_path):
    a = analyze(small_program, iterations=200)
    path = tmp_path / "s.xlsx"
    build_workbook(a, path)
    status = (date(2027, 1, 18) - date(1899, 12, 30)).days
    xl = _recalc_after_edit(path, {"Summary!B4": status}, ["Summary!B5", "EVM!O4"])
    assert xl.evaluate("Summary!B5") == 10
    # PV at 10 elapsed: A3000 + B4000 + C2000 + D5000*3/5 + E6000*5/6 = 17000
    assert xl.evaluate("EVM!O4") == pytest.approx(17000)


def test_small_workbook_fully_verifies(small_program, tmp_path):
    a = analyze(small_program, iterations=200)
    _, expected = build_workbook(a, tmp_path / "s.xlsx")
    assert verify_workbook(tmp_path / "s.xlsx", expected).ok


def test_large_network_with_gantt_cap(tmp_path):
    a = analyze(generate(500, seed=1), iterations=200)
    wb, expected = build_workbook(a, tmp_path / "big.xlsx")
    assert len(expected) > 8000
